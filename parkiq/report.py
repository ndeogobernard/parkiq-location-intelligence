"""Private partner documents: investment memo, site profiles, model placeholder, package (M7 early).

* :func:`memo_numbers` computes every number the memo quotes, each with its source (table,
  field or report key); :func:`build_memo` fills ``templates/memo.md`` and appends the number
  trace table. Financial sections stay marked **[M6]** until underwriting exists.
* :func:`build_profiles` writes one page per top-ranked site: map (map standard, private),
  facts, flags, due-diligence checklist; financials left for M6.
* :func:`build_model_placeholder` writes the Excel model's structure (inputs with their status,
  site inputs from GIS); formulas arrive in M6.
* :func:`build_package` bundles memo, profiles, model placeholder, the private dashboard and the
  run's maps into ``<run>/package/`` and a zip. Everything here is PRIVATE (candidate sites).
"""

from __future__ import annotations

import json
import shutil
import sqlite3
import zipfile
from contextlib import closing
from datetime import date
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import pyogrio

from parkiq import portfolio

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "memo.md"
EAST_OF_HIGH_ST = -82.9988  # N High St through downtown Columbus (Franklin market only)
CRIT = {
    "c01": "Weekday-day shortage nearby", "c02": "Evening and weekend shortage nearby",
    "c03": "Event shortage (neutral)", "c04": "Achievable rate (neutral)",
    "c05": "Competing supply (lower is better)", "c06": "Walk to main destinations",
    "c07": "Land cost per stall", "c08": "Access (corner, arterial, traffic)",
    "c09": "Zoning certainty", "c10": "Planned projects (neutral)",
}  # fmt: skip
PRIVATE_CSS = (
    "\nimg { max-width: 100%%; max-height: 3.9in; height: auto; display: block; margin: 0 auto; }"
    "\nsection.profile { page-break-before: always; }"
    "\nsection.profile td, section.profile th { padding: 1.5pt 5pt; font-size: 8.5pt; }"
    "\nsection.profile p, section.profile li { margin: 2pt 0; }\n"
)


def _lower_first(s: str) -> str:
    """'Near Ohio State' -> 'near Ohio State' (mid-sentence)."""
    return s[:1].lower() + s[1:]


def _fmt(v: float) -> str:
    return f"{round(v):,}"


def _json(run_dir: Path, name: str) -> dict[str, Any]:
    f = run_dir / name
    return dict(json.loads(f.read_text(encoding="utf-8"))) if f.exists() else {}


def _count(gpkg: Path, sql: str) -> int:
    with closing(sqlite3.connect(gpkg)) as con:
        return int(con.execute(sql).fetchone()[0])


def _table(df: pd.DataFrame) -> str:
    head = "| " + " | ".join(df.columns) + " |"
    sep = "|" + "---|" * len(df.columns)
    rows = ["| " + " | ".join(str(v) for v in r) + " |" for r in df.itertuples(index=False)]
    return "\n".join([head, sep, *rows])


def ranked_candidates(gpkg: Path, scenario: str = "Balanced") -> pd.DataFrame:
    """Candidates with their rank and score in one scenario, best first (with geometry)."""
    cand = pyogrio.read_dataframe(gpkg, layer="CandidateParcels", where="screen_status <> 'Fail'")
    sc = pyogrio.read_dataframe(gpkg, layer="SiteScores", read_geometry=False)
    s = sc[sc["scenario"] == scenario][["parcel_id", "rank", "composite"]]
    return cand.merge(s, on="parcel_id").sort_values("rank").reset_index(drop=True)


def memo_numbers(ctx: Any) -> dict[str, tuple[str, str]]:
    """Every memo number with its source: {key: (value, source)}."""
    from parkiq.dashboard import zone_labels

    cfg, run, gpkg = ctx.cfg, ctx.run_dir, ctx.store.gpkg
    dem, sup, gap = (
        _json(run, "demand_report.json"),
        _json(run, "supply_report.json"),
        _json(run, "gap_report.json"),
    )
    scr, sqlr = _json(run, "screen_report.json"), _json(run, "sql_check_report.json")
    n: dict[str, tuple[str, str]] = {}
    n["market_name"] = (cfg.market.market.name, "markets/<slug>.yaml market.name")
    n["run_id"] = (run.name, "run folder")
    n["analysis_month"] = (date.today().strftime("%B %Y"), "build date")
    n["stalls_min"] = (str(cfg.market.site.target_stalls_range[0]), "site.target_stalls_range")
    n["stalls_max"] = (str(cfg.market.site.target_stalls_range[1]), "site.target_stalls_range")
    n["parcels"] = (_fmt(_count(gpkg, 'SELECT COUNT(*) FROM "Parcels"')), "COUNT(Parcels)")
    n["hexes"] = (_fmt(_count(gpkg, 'SELECT COUNT(*) FROM "HexGrid"')), "COUNT(HexGrid)")
    n["sources_loaded"] = (_fmt(_count(gpkg, "SELECT COUNT(*) FROM DataSourceRegistry WHERE status = 'loaded'")),
                           "DataSourceRegistry.status = loaded")  # fmt: skip
    n["qa_checks"] = (_fmt(_count(gpkg, 'SELECT COUNT(*) FROM "QAQC_Log"')), "COUNT(QAQC_Log)")
    n["qa_errors"] = (_fmt(_count(gpkg, "SELECT COUNT(*) FROM QAQC_Log WHERE passed = 0 AND severity = 'error'")),
                      "QAQC_Log failed error-level checks")  # fmt: skip
    n["sql_checks"] = (_fmt(sqlr.get("checks", 0)), "sql_check_report.json checks")
    n["sql_failed"] = (_fmt(len(sqlr.get("failed", []))), "sql_check_report.json failed")
    n["anchors"] = (
        _fmt(_count(gpkg, 'SELECT COUNT(*) FROM "DemandAnchors"')),
        "COUNT(DemandAnchors)",
    )
    for dp in ("wd_day", "wd_eve", "we_day", "we_eve"):
        n[f"demand_{dp}"] = (
            _fmt(dem.get("demand_total", {}).get(dp, 0)),
            f"demand_report.json demand_total.{dp}",
        )
    att = cfg.market.demand.attendance_factor.get("Office", {}).get("wd_day")
    n["attendance_factor"] = (
        f"{att:.2f}" if att else "not set",
        "demand.attendance_factor.Office.wd_day",
    )
    n["mapped_facilities"] = (
        _fmt(_count(gpkg, "SELECT COUNT(*) FROM SupplyFacilities WHERE capacity_source IS NULL "
                          "OR capacity_source <> 'parcel estimate'")),
        "SupplyFacilities excluding parcel estimates",
    )  # fmt: skip
    n["metered_faces"] = (
        _fmt(_count(gpkg, 'SELECT COUNT(*) FROM "OnStreetSegments"')),
        "COUNT(OnStreetSegments)",
    )
    pe = sup.get("parcel_estimate", {})
    n["estimated_parcels"] = (
        _fmt(pe.get("parcels_estimated", 0)),
        "supply_report.json parcel_estimate",
    )
    n["estimated_stalls"] = (_fmt(pe.get("stalls", 0)), "supply_report.json parcel_estimate.stalls")
    n["paid_hexes"] = (
        _fmt(gap.get("paid_market", {}).get("paid_market_hexes", 0)),
        "gap_report.json paid_market",
    )
    n["paid_zones"] = (
        _fmt(gap.get("screening", {}).get("zones", 0)),
        "gap_report.json screening.zones",
    )
    n["paid_zones_50"] = (
        _fmt(gap.get("screening", {}).get("zones_ge_min_lot", 0)),
        "gap_report.json screening",
    )
    n["context_zones"] = (
        _fmt(gap.get("context_only", {}).get("zones", 0)),
        "gap_report.json context_only",
    )
    n["within_reach"] = (
        _fmt(scr.get("within_reach_sites", 0)),
        "screen_report.json within_reach_sites",
    )
    n["candidates"] = (_fmt(scr.get("candidates", 0)), "screen_report.json candidates")
    byst = scr.get("candidates_by_status", {})
    n["cand_pass"] = (_fmt(byst.get("Pass", 0)), "screen_report.json candidates_by_status.Pass")
    n["cand_review"] = (
        _fmt(byst.get("Review", 0)),
        "screen_report.json candidates_by_status.Review",
    )
    cands = ranked_candidates(gpkg)
    n["assemblies"] = (
        _fmt(int((cands["assembled_count"] > 1).sum())),
        "CandidateParcels.assembled_count > 1",
    )
    n["existing_only"] = (
        _fmt(int(cands["existing_only_flag"].fillna(0).astype(bool).sum())),
        "CandidateParcels.existing_only_flag",
    )
    n["zone_a"] = (
        _fmt(int(cands["zone_a_flag"].fillna(0).astype(bool).sum())),
        "CandidateParcels.zone_a_flag",
    )
    n["owner_flagged"] = (
        _fmt(int(cands["owner_feasibility"].notna().sum())),
        "CandidateParcels.owner_feasibility",
    )
    lon = cands.head(10).geometry.representative_point().to_crs(4326).x
    n["top10_east"] = (
        _fmt(int((lon > EAST_OF_HIGH_ST).sum())),
        "SiteScores Balanced rank 1-10, longitude > N High St",
    )
    hz = pyogrio.read_dataframe(gpkg, layer="HotZones")
    paid = hz[hz["zone_class"].astype(str).str.startswith("paid")].sort_values(
        "total_gap_stalls", ascending=False
    )
    if len(paid):
        subs = pyogrio.read_dataframe(gpkg, layer="Submarkets")
        subs["geometry"] = subs.geometry.make_valid()
        anchors = pyogrio.read_dataframe(gpkg, layer="DemandAnchors", read_geometry=False)
        top = paid.head(1)
        n["largest_zone_name"] = (_lower_first(zone_labels(top, {"anchors": anchors, "submarkets": subs})[0]),
                                  "HotZones (paid) max total_gap_stalls")  # fmt: skip
        n["largest_zone_gap"] = (
            _fmt(float(top["total_gap_stalls"].iloc[0])),
            "HotZones.total_gap_stalls",
        )
        zid = str(top["zone_id"].iloc[0])
        n["largest_zone_candidates"] = (
            _fmt(int((cands["hotzone_id"] == zid).sum())),
            "CandidateParcels.hotzone_id",
        )
    top10 = cands.head(10)
    t = pd.DataFrame({
        "Rank": top10["rank"].astype(int), "Address": top10["address"].fillna("no site address"),
        "Stalls": top10["stalls"].map(_fmt), "Zoning": top10["zoning_screen"],
        "Screen": top10["screen_status"], "Score": top10["composite"].round(1),
        "Owner": top10["owner_feasibility"].fillna(top10["owner_type"]),
    })  # fmt: skip
    n["top10_table"] = (_table(t), "SiteScores (Balanced rank, composite), CandidateParcels")
    bench = cfg.market.supply.benchmark_occupancy_path
    if bench and Path(bench).exists():
        b = pd.read_csv(bench)
        pc = cfg.market.supply.practical_capacity or 0.85
        bt = pd.DataFrame({
            "Area": b["area"], "Surveyed spaces": b["surveyed_spaces"].map(_fmt),
            "Peak occupancy": b["peak_occupancy"].map(lambda v: f"{v:.1%}"),
            "vs 85 %": b["peak_occupancy"].map(lambda v: "above" if v >= pc else "below"),
        })  # fmt: skip
        n["benchmark_table"] = (_table(bt), "markets/<slug>/benchmark_occupancy.csv")
    else:
        n["benchmark_table"] = ("no benchmark file", "supply.benchmark_occupancy_path")
    reg = pyogrio.read_dataframe(gpkg, layer="DataSourceRegistry", read_geometry=False)
    n["sources_table"] = (portfolio.sources_table(reg), "DataSourceRegistry")
    site = cfg.market.site
    params = pd.DataFrame({
        "Parameter": ["Lot size (sq ft)", "Stalls", "Stall area (sq ft, gross)", "Layout efficiency",
                      "Minimum shape index", "Minimum frontage (ft)", "Maximum slope (%)",
                      "Walk to paid-parking hot zone (min)"],
        "Value": [f"{site.min_parcel_sqft:,.0f} to {site.max_parcel_sqft:,.0f}",
                  f"{site.target_stalls_range[0]} to {site.target_stalls_range[1]}",
                  f"{site.stall_area_sqft_gross:g}", f"{site.layout_efficiency:g}", f"{site.min_shape_index:g}",
                  f"{site.min_frontage_ft:g}", f"{site.max_slope_pct:g}", "8"],
    })  # fmt: skip
    n["parameters_table"] = (_table(params), "params.yaml market.site")
    return n


def build_memo(ctx: Any, out_dir: Path, template: Path = TEMPLATE) -> Path:
    """Fill the memo template; append the number trace; write .md and .pdf."""
    nums = memo_numbers(ctx)
    used = sorted(
        set(portfolio.PLACEHOLDER.findall(template.read_text(encoding="utf-8"))) - {"trace_table"}
    )
    trace = pd.DataFrame(
        [{"Number": k, "Value": nums[k][0] if "\n" not in nums[k][0] else "(table)", "Source": nums[k][1]}
         for k in used if k in nums]
    )  # fmt: skip
    values = {k: v for k, (v, _) in nums.items()} | {"trace_table": _table(trace)}
    out_dir.mkdir(parents=True, exist_ok=True)
    src = out_dir / "ParkIQ_investment_memo.md"
    shutil.copy(template, src)
    pdf = portfolio.render(src, out_dir / "ParkIQ_investment_memo.pdf", values, "Investment memo (PRIVATE)",
                           draft=True, links=False, draft_label="PRIVATE DRAFT")  # fmt: skip
    filled = portfolio.PLACEHOLDER.sub(
        lambda m: values[m.group(1)], src.read_text(encoding="utf-8")
    )
    src.write_text(filled, encoding="utf-8")
    return pdf


def _site_map(
    gpkg: Path, c: Any, pid: str, k: int, sheds: Any, layers: dict[str, Any], out: Path
) -> Path:
    import shapely
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    from parkiq.maps import MapSpec, finish, locator, new_map, north_arrow, scale_bar

    county, hz, pz, ons, fac, cand = (
        layers[x] for x in ("county", "hz", "pz", "ons", "fac", "cand")
    )
    view = tuple(np.array(c.geometry.bounds) + np.array([-2600, -2000, 2600, 2000]))
    box = shapely.geometry.box(*view)
    edges = pyogrio.read_dataframe(gpkg, layer="WalkEdges", columns=["highway"], bbox=view)
    spec = MapSpec(
        title=f"Site {k}: {c.address or 'no site address'}",
        subtitle="Candidate site, 3/5/8-minute walk sheds, parking nearby and hot zones, weekday",
        how_to_read=("The red outline is the site. Shaded rings show where you can walk in 3, 5 and 8 "
                     "minutes. Dots are lots (blue) and garages (orange); green lines are metered streets."),
        sources=["County parcels (run)", "OSM parking & streets (run)", "City curb meters (run)",
                 "ParkIQ walk sheds and hot zones (run)"],
        uses_osm=True, result=True, public=False, units="stalls", market_extent=tuple(county.total_bounds),
    )  # fmt: skip
    fig, ax = new_map(spec, figsize=(14, 8.4))
    edges.plot(ax=ax, color="#d9d9d9", lw=0.5)
    s = sheds[sheds["parcel_id"] == pid].sort_values("minutes", ascending=False)
    for (_, r), col in zip(s.iterrows(), ["#c6dbef", "#6baed6", "#2171b5"], strict=False):
        s[s["minutes"] == r["minutes"]].plot(ax=ax, color=col, alpha=0.35, lw=0)
    hz[hz.intersects(box) & hz["zone_class"].astype(str).str.startswith("paid")].boundary.plot(
        ax=ax, color="black", lw=1.6
    )
    pzd = pz.dissolve(by="zone").reset_index()
    for zone, col in (("A", "#fde0dd"), ("B", "#fff7bc")):
        pzd[(pzd["zone"] == zone) & pzd.intersects(box)].plot(
            ax=ax, color=col, alpha=0.5, lw=0, zorder=0
        )
    ons[ons.intersects(box)].plot(ax=ax, color="#1b9e77", lw=1.2)
    f = fac[fac.intersects(box)]
    cap = f["capacity_stated"].fillna(f["capacity_est"]).fillna(0)
    for typ, col in (("Surface", "#2c7fb8"), ("Garage", "#d95f0e")):
        m = f["type"] == typ
        pt = f[m].geometry.representative_point()
        ax.scatter(
            pt.x,
            pt.y,
            s=np.clip(cap[m] / 5, 6, 400),
            color=col,
            alpha=0.7,
            edgecolor="white",
            lw=0.4,
            zorder=4,
        )
    cand.loc[[pid]].boundary.plot(ax=ax, color="#e31a1c", lw=2.5, zorder=6)
    ax.set_xlim(view[0], view[2])
    ax.set_ylim(view[1], view[3])
    scale_bar(ax, 1320, "¼ mi")
    north_arrow(ax)
    locator(fig, county, view)
    leg = [
        Line2D([0], [0], color="#e31a1c", lw=2.5, label="candidate site"),
        Patch(color="#2171b5", alpha=0.35, label="3-minute walk"), Patch(color="#6baed6", alpha=0.35, label="5-minute walk"),
        Patch(color="#c6dbef", alpha=0.35, label="8-minute walk"),
        Line2D([0], [0], marker="o", color="w", markerfacecolor="#2c7fb8", markersize=9, label="surface lot (sized by stalls)"),
        Line2D([0], [0], marker="o", color="w", markerfacecolor="#d95f0e", markersize=9, label="garage"),
        Line2D([0], [0], color="#1b9e77", lw=2, label="metered street"),
        Line2D([0], [0], color="black", lw=2, label="paid-parking hot zone"),
        Patch(color="#fde0dd", label="Downtown parking Zone A (no new surface lots)"),
        Patch(color="#fff7bc", label="Downtown parking Zone B"),
    ]  # fmt: skip
    path = out / f"site-{k:02d}.png"
    finish(fig, ax, leg, "Map", path, view=view)
    return path


def checklist(c: Any) -> list[str]:
    """Rule-based due-diligence items for one candidate (ADR-0046)."""
    items = [
        f"Confirm zoning and any required approval with the city ({c.zoning_code}; screen: {c.zoning_screen})",
        "Title search, owner contact and willingness to sell or lease",
        "Boundary and ALTA survey; confirm lot size and buildable stalls",
        "Phase I environmental site assessment",
    ]
    if c.existing_lot_flag:
        items.append(
            "Existing lot: verify legal nonconforming status and continuous use; inspect pavement and lighting"
        )
    if c.zone_a_flag:
        items.append(
            "Downtown parking Zone A: confirm the Zone A/B line on the official map; new surface lots are not allowed in Zone A"
        )
    if "near Zone A/B" in str(c.screen_reason):
        items.append("Site is near the digitized Zone A/B line: confirm which zone applies")
    if isinstance(c.owner_feasibility, str):
        items.append(f"Owner: {c.owner_feasibility}; open a ground lease discussion")
    if c.flood_flag:
        items.append("In a FEMA special flood hazard area: floodplain permit and insurance")
    if c.arterial_flag:
        items.append("Curb cut on an arterial street: access approval from the city")
    if c.assembled_count and c.assembled_count > 1:
        items.append(
            f"Assembly of {c.assembled_count} parcels: all must be acquired or leased together"
        )
    items.append("Landscaping, screening and stormwater requirements for parking lots")
    return items


def build_profiles(ctx: Any, out_dir: Path, n: int = 10) -> Path:
    """One private page per top-``n`` site (Balanced rank)."""
    gpkg = ctx.store.gpkg
    maps = out_dir / "profile_maps"
    maps.mkdir(parents=True, exist_ok=True)
    cands = ranked_candidates(gpkg)
    sc = pyogrio.read_dataframe(gpkg, layer="SiteScores", read_geometry=False)
    rs = pyogrio.read_dataframe(gpkg, layer="RankStability", read_geometry=False).set_index(
        "parcel_id"
    )
    fac = pyogrio.read_dataframe(gpkg, layer="SupplyFacilities")
    layers = {
        "county": pyogrio.read_dataframe(gpkg, layer="MarketBoundary"),
        "hz": pyogrio.read_dataframe(gpkg, layer="HotZones"),
        "pz": pyogrio.read_dataframe(gpkg, layer="ParkingZones"),
        "ons": pyogrio.read_dataframe(gpkg, layer="OnStreetSegments"),
        "fac": fac[fac["capacity_source"] != "parcel estimate"],
        "cand": cands.set_index("parcel_id"),
    }
    sheds = pyogrio.read_dataframe(gpkg, layer="WalkSheds")
    pages = []
    for k, (_, c) in enumerate(cands.head(n).iterrows(), 1):
        pid = c["parcel_id"]
        ranks = sc[sc["parcel_id"] == pid].set_index("scenario")
        st = rs.loc[pid]
        mp = _site_map(gpkg, c, pid, k, sheds, layers, maps)
        r = sc[(sc["parcel_id"] == pid) & (sc["scenario"] == "Balanced")].iloc[0]
        flags = [
            x.strip() for x in str(c.screen_reason).split(";") if x.strip() and x.strip() != "PASS"
        ]
        where = ("inside a paid-parking hot zone" if c.walk_min_to_hotzone == 0
                 else f"{c.walk_min_to_hotzone:.1f} min walk to a paid-parking hot zone")  # fmt: skip
        facts = [
            ("Location", f"{c.address or 'no site address'}, {c.jurisdiction}; {where}"),
            ("Lot", f"{c.lot_sqft:,.0f} sq ft ({c.lot_sqft / 43560:.2f} acres); {c.assembled_count} parcel(s)"),
            ("Stalls", f"{c.stalls:,.0f} ({c.stalls_basis})"),
            ("Zoning", f"{c.zoning_code}: {c.zoning_screen}; screen {c.screen_status}"),
            ("Existing use", ("surface lot, " if c.existing_lot_flag else "") + str(c.land_use_class)),
            ("Owner", f"{c.owner_type}" + (f"; {c.owner_feasibility}" if isinstance(c.owner_feasibility, str) else "")),
            ("Access", f"frontage {c.frontage_ft:,.0f} ft; corner {'yes' if c.corner_flag else 'no'}; arterial {'yes' if c.arterial_flag else 'no'}"),
            ("Site conditions", f"slope {c.slope_pct:.1f}%; flood zone {'yes' if c.flood_flag else 'no'}; brownfield {'yes' if c.brownfield_flag else 'no'}"),
            ("Rank", " / ".join(f"{s}: {int(ranks.loc[s, 'rank'])} (score {ranks.loc[s, 'composite']:.1f})"
                                for s in ranks.index)),
            ("Stability", f"top 10 in {st.top10_freq:.0%} of 1,000 random weightings; median rank {st.median_rank:.0f}"),
        ]  # fmt: skip
        crit = " · ".join(f"{CRIT[x]} {r[f'{x}_s']:.0f}" for x in CRIT)
        body = [
            f"![Site {k} map]({maps.name}/{mp.name})", "",
            "| | |", "|---|---|", *[f"| **{a}** | {b} |" for a, b in facts], "",
            f"**Criteria scores (0 to 100, higher is better):** {crit}", "",
            "**Flags:** " + ("; ".join(dict.fromkeys(flags)) if flags else "none"), "",
            "**Due diligence:**", "", *[f"- [ ] {x}" for x in checklist(c)], "",
            "**Financials (M6):** buy and ground-lease underwriting follows once the partners set the finance assumptions.",
        ]  # fmt: skip
        pages.append(
            '<section markdown="1" class="profile">\n\n' + "\n".join(body) + "\n\n</section>"
        )
    md = (f"# Top {n} candidate sites: PRIVATE site profiles\n\nFor partner preview only; not for distribution. "
          f"Ranks can change once rates, venues and planned projects are added. Run {ctx.run_dir.name}.\n\n"
          + "\n".join(pages))  # fmt: skip
    src = out_dir / "ParkIQ_site_profiles.md"
    src.write_text(md, encoding="utf-8")
    return portfolio.render(src, out_dir / "ParkIQ_site_profiles.pdf", {}, "Site profiles (PRIVATE)",
                            draft=True, extra_css=PRIVATE_CSS, links=False, draft_label="PRIVATE")  # fmt: skip


def build_model_placeholder(ctx: Any, out: Path, n: int = 10) -> Path:
    """Excel model structure until M6: finance inputs with status, site inputs from GIS."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fin = ctx.cfg.finance.model_dump()
    prov = ctx.cfg.provenance
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(
        [
            "PLACEHOLDER: structure only. Live formulas arrive in M6; blank = partner decision (DECIDE)."
        ]
    )
    ws["A1"].font = Font(bold=True, color="C00000")
    ws.append(["Parameter", "Value", "Status", "Source"])
    for k, v in fin.items():
        p = prov.get(f"finance.{k}")
        status = p.status if p else ("DECIDE" if v is None else "")
        ws.append(
            [
                k,
                "" if v is None else (json.dumps(v) if isinstance(v, dict | list) else v),
                status,
                p.source if p else "",
            ]
        )
        if v is None:
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor="FFF2CC")
    sites = wb.create_sheet("Sites")
    cands = ranked_candidates(ctx.store.gpkg).head(n)
    cols = ["rank", "parcel_id", "address", "lot_sqft", "stalls", "stalls_basis", "land_value", "zoning_screen",
            "existing_lot_flag", "owner_feasibility"]  # fmt: skip
    sites.append(cols)
    for _, r in cands.iterrows():
        sites.append(
            [
                None if pd.isna(r[c]) else (r[c].item() if hasattr(r[c], "item") else r[c])
                for c in cols
            ]
        )
    for name in ("Occupancy", "CF_Buy", "CF_Lease", "Sensitivity", "Summary"):
        s = wb.create_sheet(name)
        s.append([f"{name}: M6 (after the partners set the finance assumptions)"])
    wb.save(out)
    return out


def build_package(ctx: Any, maps_dirs: list[Path] | None = None, n_profiles: int = 10) -> Path:
    """PRIVATE partner package in <run>/package/ plus a zip next to it."""
    from parkiq import dashboard

    pkg = ctx.run_dir / "package"
    if pkg.exists():
        shutil.rmtree(pkg)
    pkg.mkdir(parents=True)
    build_memo(ctx, pkg)
    build_profiles(ctx, pkg, n_profiles)
    build_model_placeholder(ctx, pkg / "ParkIQ_financial_model_PLACEHOLDER.xlsx", n_profiles)
    dashboard.build(ctx.store.gpkg, pkg / "ParkIQ_dashboard_PRIVATE.html", public=False,
                    analysis_date=date.today().strftime("%B %Y"))  # fmt: skip
    mdir = pkg / "maps"
    mdir.mkdir()
    for d in [ctx.run_dir / "maps", *(maps_dirs or [])]:
        if d.exists():
            for f in sorted(d.glob("*")):
                if f.suffix.lower() in (".png", ".jpg", ".jpeg"):
                    shutil.copy(f, mdir / f.name)
    files = sorted(p.relative_to(pkg).as_posix() for p in pkg.rglob("*") if p.is_file())
    (pkg / "README.txt").write_text(
        "ParkIQ partner package (PRIVATE: contains candidate sites; do not publish or forward)\n"
        f"Run {ctx.run_dir.name}, built {date.today().isoformat()}\n\n" + "\n".join(files) + "\n",
        encoding="utf-8",
    )
    zpath = Path(ctx.run_dir) / f"ParkIQ_partner_package_{ctx.run_dir.name}.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(pkg.rglob("*")):
            if p.is_file():
                z.write(p, p.relative_to(pkg).as_posix())
    return zpath
